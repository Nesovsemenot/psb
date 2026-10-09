#!/usr/bin/env python3
"""
Отчёт по статьям psbank.ru/articles из Яндекс.Метрики.
Период: июнь 2025 — сентябрь 2026, помесячно.

Запуск:
  python3 ym_articles_report.py --auth       ← получить OAuth-токен (один раз)
  python3 ym_articles_report.py              ← выгрузить данные → xlsx
  python3 ym_articles_report.py --diag-ai    ← показать все источники ИИ-трафика
"""

import argparse
import calendar
import json
import sys
import time
import webbrowser
from datetime import datetime
from pathlib import Path

import requests
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ═══════════════════════════════════════════════════════════
# Конфиг
# ═══════════════════════════════════════════════════════════

COUNTER   = 18882175
CLIENT_ID = "b57a63b5c3f442409b3e8a033d31fb4e"

GOAL_APP  = 465979958   # «Все статьи // Успешная отправка заявки // 12.09.2025»
GOAL_CTA  = 519026502   # «МСП // Статьи // CTA»

DATE1 = "2025-06-01"
DATE2 = "2026-09-30"

API = "https://api-metrika.yandex.net/stat/v1/data"
TOKEN_FILE = Path(__file__).with_name(".ym_token")

# ИИ-источники: домены рефереров
AI_DOMAINS = [
    # ChatGPT
    "chatgpt.com", "chat.openai.com",
    # Perplexity
    "perplexity.ai",
    # Claude
    "claude.ai",
    # Gemini
    "gemini.google.com", "business.gemini.google",
    # Copilot
    "copilot.microsoft.com", "copilot.com", "copilot.cloud.microsoft",
    # DeepSeek
    "deepseek.com",
    # Grok
    "grok.com", "x.ai",
    # Qwen
    "qwen.ai",
    # Kimi
    "kimi.com", "kimi.moonshot.cn",
    # Mistral
    "mistral.ai",
    # Алиса
    "alice.yandex.ru", "alicepro.yandex.ru", "alice.ya.ru",
    # GigaChat
    "giga.chat", "gigachat.ru", "gigachat.devices.sberbank.ru",
    # Poe
    "poe.com",
    # Phind
    "phind.com",
    # You.com
    "you.com",
    # Meta AI
    "meta.ai",
    # Duck.ai
    "duck.ai",
    # Genspark
    "genspark.ai",
    # Felo
    "felo.ai",
]

# ИИ-источники: значения utm_source
AI_UTM_SOURCES = [
    "chatgpt", "chatgpt.com", "openai",
    "perplexity", "perplexity.ai",
    "claude", "claude.ai",
    "gemini", "gemini.google.com",
    "copilot", "copilot.com",
    "deepseek",
    "grok",
    "qwen",
    "kimi",
    "mistral",
    "alice",
    "gigachat", "giga.chat",
    "poe",
    "phind",
]

# Паттерны для определения сегмента
MSB_PATTERNS   = ["/msb", "/business", "/predprinimatelyam"]
RETAIL_PATTERNS = ["/rb", "/retail", "/chastnye", "/fizicheskim"]

PAUSE = 0.35  # пауза между запросами


def normalize_path(p: str) -> str:
    """Обрезает query-параметры: всё после первого ? или & и trailing /."""
    for ch in ("?", "&"):
        idx = p.find(ch)
        if idx != -1:
            p = p[:idx]
    return p.rstrip("/")


# ═══════════════════════════════════════════════════════════
# Утилиты
# ═══════════════════════════════════════════════════════════

def load_token() -> str:
    if TOKEN_FILE.exists():
        return TOKEN_FILE.read_text().strip()
    print(f"Токен не найден ({TOKEN_FILE}).")
    print("Запустите:  python3 ym_articles_report.py --auth")
    sys.exit(1)


def save_token(token: str):
    TOKEN_FILE.write_text(token)
    print(f"Токен сохранён → {TOKEN_FILE}")


def do_auth():
    url = (
        f"https://oauth.yandex.ru/authorize?response_type=token"
        f"&client_id={CLIENT_ID}"
    )
    print("Открываю браузер для авторизации…")
    print(f"Если не открылся, перейдите вручную:\n{url}\n")
    webbrowser.open(url)
    print("После авторизации Яндекс перенаправит на страницу с токеном в URL.")
    print("Скопируйте значение access_token и вставьте сюда:\n")
    token = input("OAuth-токен: ").strip()
    if not token:
        print("Пустой токен, отмена.")
        sys.exit(1)
    save_token(token)
    print("Готово. Теперь запустите скрипт без --auth.")


def detect_segment(path: str) -> str:
    lp = path.lower()
    for p in MSB_PATTERNS:
        if p in lp:
            return "МСП"
    for p in RETAIL_PATTERNS:
        if p in lp:
            return "Розница"
    return ""


def month_range(y: int, m: int):
    """Возвращает ('YYYY-MM-DD', 'YYYY-MM-DD') для начала и конца месяца."""
    last_day = calendar.monthrange(y, m)[1]
    return f"{y}-{m:02d}-01", f"{y}-{m:02d}-{last_day:02d}"


def months_list():
    """Список кортежей (year, month, 'YYYY-MM') за весь период."""
    result = []
    y, m = 2025, 6
    while (y, m) <= (2026, 9):
        result.append((y, m, f"{y}-{m:02d}"))
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return result


# ═══════════════════════════════════════════════════════════
# API
# ═══════════════════════════════════════════════════════════

def api_get(token: str, params: dict) -> dict:
    headers = {"Authorization": f"OAuth {token}"}
    params["ids"]  = COUNTER
    params["lang"] = "ru"
    params["accuracy"] = 1
    params.setdefault("limit", 10000)
    params.setdefault("offset", 1)

    all_data = []

    while True:
        for attempt in range(5):
            try:
                resp = requests.get(API, params=params, headers=headers, timeout=300)
            except requests.exceptions.ReadTimeout:
                print(f" T/O({attempt+1})", end="", flush=True)
                time.sleep(5)
                continue
            if resp.status_code == 429:
                wait = 5 * (attempt + 1)
                print(f" RL({wait}s)", end="", flush=True)
                time.sleep(wait)
                continue
            break
        else:
            print("\n  ❌ Все 5 попыток неудачны")
            sys.exit(1)
        if not resp.ok:
            print(f"\n  ❌ API {resp.status_code}: {resp.text[:300]}")
            resp.raise_for_status()
        body = resp.json()
        all_data.extend(body.get("data", []))

        total = body.get("total_rows", 0)
        fetched = params["offset"] - 1 + len(body.get("data", []))
        if fetched >= total:
            break
        params["offset"] = fetched + 1
        time.sleep(PAUSE)

    return all_data


def fetch_one_month(token: str, d1: str, d2: str,
                    extra_filter: str = "",
                    metrics: str = "ym:s:visits") -> dict[str, float]:
    """Возвращает {url_path: value} за один месяц."""
    base = "ym:s:startURLPath=@'/articles/'"
    filt = f"{base} AND {extra_filter}" if extra_filter else base

    params = {
        "metrics":    metrics,
        "dimensions": "ym:s:startURLPath",
        "date1":      d1,
        "date2":      d2,
        "filters":    filt,
    }
    rows = api_get(token, params)
    result = {}
    for row in rows:
        raw_path = row["dimensions"][0].get("name", "")
        val = row["metrics"][0] if row["metrics"] else 0
        path = normalize_path(raw_path)
        if path:
            result[path] = result.get(path, 0) + int(round(val))
    return result


def fetch_titles(token: str) -> dict[str, str]:
    """URL path → title через pageviews API."""
    params = {
        "metrics":    "ym:pv:pageviews",
        "dimensions": "ym:pv:URLPath,ym:pv:title",
        "date1":      DATE1,
        "date2":      DATE2,
        "filters":    "ym:pv:URLPath=@'/articles/'",
    }
    rows = api_get(token, params)
    titles = {}
    for row in rows:
        raw   = row["dimensions"][0].get("name", "")
        title = row["dimensions"][1].get("name") or ""
        path  = normalize_path(raw)
        if title in ("", "Страница не найдена"):
            continue
        if path and (path not in titles or len(title) > len(titles.get(path, ""))):
            titles[path] = title
    return titles


def ai_filter_chunks(max_conds: int = 18) -> list[str]:
    """Разбивает ИИ-фильтр на куски по max_conds условий (лимит API — 20,
    минус 1 на URL-фильтр, минус 1 запас)."""
    all_parts = [f"ym:s:referer=@'{d}'" for d in AI_DOMAINS]
    all_parts += [f"ym:s:lastUTMSource=='{s}'" for s in AI_UTM_SOURCES]
    chunks = []
    for i in range(0, len(all_parts), max_conds):
        chunk = all_parts[i : i + max_conds]
        chunks.append("(" + " OR ".join(chunk) + ")")
    return chunks


# ═══════════════════════════════════════════════════════════
# Диагностика ИИ-источников
# ═══════════════════════════════════════════════════════════

def diag_ai(token: str):
    """Показать все поисковые системы и рефереры для /articles/."""
    print("\n══ Диагностика ИИ-источников на /articles/ ══\n")

    # 1. Поисковые системы
    print("① Поисковые системы (топ-30):")
    rows = api_get(token, {
        "metrics": "ym:s:visits",
        "dimensions": "ym:s:lastSearchEngine",
        "date1": DATE1, "date2": DATE2,
        "filters": "ym:s:startURLPath=@'/articles/'",
        "sort": "-ym:s:visits",
        "limit": 30,
    })
    for r in rows:
        name = r["dimensions"][0].get("name", "?")
        did  = r["dimensions"][0].get("id", "")
        vis  = int(r["metrics"][0])
        print(f"   {name:<40s} (id={did})  {vis:>8,} визитов")

    # 2. Топ рефереров (не поиск)
    print("\n② Реферальные источники (топ-30):")
    rows = api_get(token, {
        "metrics": "ym:s:visits",
        "dimensions": "ym:s:lastReferalSource",
        "date1": DATE1, "date2": DATE2,
        "filters": "ym:s:startURLPath=@'/articles/' AND ym:s:lastTrafficSource=='referral'",
        "sort": "-ym:s:visits",
        "limit": 30,
    })
    for r in rows:
        name = r["dimensions"][0].get("name", "?")
        vis  = int(r["metrics"][0])
        print(f"   {name:<50s}  {vis:>8,} визитов")

    # 3. Все источники трафика (сводка)
    print("\n③ Типы трафика (сводка):")
    rows = api_get(token, {
        "metrics": "ym:s:visits",
        "dimensions": "ym:s:lastTrafficSource",
        "date1": DATE1, "date2": DATE2,
        "filters": "ym:s:startURLPath=@'/articles/'",
        "sort": "-ym:s:visits",
    })
    for r in rows:
        name = r["dimensions"][0].get("name", "?")
        vis  = int(r["metrics"][0])
        print(f"   {name:<30s}  {vis:>10,} визитов")

    print("\n═══════════════════════════════════════════")
    print("По результатам обновите AI_DOMAINS в скрипте.")


# ═══════════════════════════════════════════════════════════
# Построение xlsx
# ═══════════════════════════════════════════════════════════

THIN   = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
H_FILL = PatternFill("solid", fgColor="2B579A")
H_FONT = Font(bold=True, color="FFFFFF", size=10)
D_FONT = Font(size=10)

METRIC_NAMES = ["Органика", "ИИ-источники", "CTA (продукт. контур)", "Заявки"]


def build_xlsx(pages, titles, data, months, out_path):
    """
    data = {
      'organic':  {(path, 'YYYY-MM'): val, …},
      'ai':       …,
      'cta':      …,
      'apps':     …,
    }
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"

    # ── Заголовки ──
    static = ["URL", "Заголовок", "Сегмент"]
    headers = list(static)
    for _, _, ym in months:
        for mn in METRIC_NAMES:
            headers.append(f"{ym}\n{mn}")

    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = BORDER
    ws.row_dimensions[1].height = 38

    # ── Данные ──
    keys = ["organic", "ai", "cta", "apps"]
    for ri, path in enumerate(pages, 2):
        ws.cell(ri, 1, f"https://www.psbank.ru{path}").font = D_FONT
        ws.cell(ri, 2, titles.get(path, "")).font = D_FONT
        ws.cell(ri, 3, detect_segment(path)).font = D_FONT
        for ci in range(1, 4):
            ws.cell(ri, ci).border = BORDER

        for mi, (_, _, ym) in enumerate(months):
            col = len(static) + mi * len(METRIC_NAMES) + 1
            for off, k in enumerate(keys):
                v = data[k].get((path, ym), 0)
                c = ws.cell(ri, col + off, value=v or 0)
                c.font = D_FONT
                c.border = BORDER
                c.alignment = Alignment(horizontal="right")

    # ── Ширины ──
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 12
    for ci in range(len(static) + 1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 13
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Мета ──
    ws2 = wb.create_sheet("Параметры прогона")
    meta = [
        ("Счётчик", COUNTER),
        ("Период", f"{DATE1} — {DATE2}"),
        ("Фильтр URL", "startURLPath содержит '/articles/'"),
        ("Цель «Заявки»", f"ID {GOAL_APP}"),
        ("Цель «CTA»", f"ID {GOAL_CTA}"),
        ("ИИ-домены", ", ".join(AI_DOMAINS)),
        ("Органика", "lastTrafficSource == 'organic'"),
        ("Дата выгрузки", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Страниц", len(pages)),
    ]
    for r, (k, v) in enumerate(meta, 1):
        ws2.cell(r, 1, k).font = Font(bold=True, size=10)
        ws2.cell(r, 2, str(v)).font = D_FONT
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 70

    wb.save(out_path)
    print(f"\n✅ Сохранено → {out_path}")
    print(f"   Страниц: {len(pages)}, месяцев: {len(months)}")


# ═══════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--auth",    action="store_true", help="Получить OAuth-токен")
    parser.add_argument("--diag-ai", action="store_true", help="Показать ИИ-источники")
    args = parser.parse_args()

    if args.auth:
        do_auth()
        return

    token = load_token()

    if args.diag_ai:
        diag_ai(token)
        return

    months = months_list()
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out_path = Path(__file__).with_name(f"articles_report_{ts}.xlsx")

    print("═" * 60)
    print("Выгрузка данных по /articles/ из Яндекс.Метрики")
    print(f"Период: {DATE1} — {DATE2}  ({len(months)} мес.)")
    print("═" * 60)

    # 0. Заголовки
    print("\n⓪ Заголовки страниц…")
    titles = fetch_titles(token)
    print(f"   Найдено: {len(titles)} стр.")
    time.sleep(PAUSE)

    # 1–4. Помесячные запросы
    data = {"organic": {}, "ai": {}, "cta": {}, "apps": {}}
    simple_queries = [
        ("organic", "① Органика",  "ym:s:lastTrafficSource=='organic'",  "ym:s:visits"),
        ("cta",     "③ CTA",       "",                                   f"ym:s:goal{GOAL_CTA}reaches"),
        ("apps",    "④ Заявки",    "",                                   f"ym:s:goal{GOAL_APP}reaches"),
    ]

    for key, label, filt, metric in simple_queries:
        print(f"\n{label}  ", end="", flush=True)
        for y, m, ym in months:
            d1, d2 = month_range(y, m)
            month_data = fetch_one_month(token, d1, d2, filt, metric)
            for path, val in month_data.items():
                data[key][(path, ym)] = val
            print("·", end="", flush=True)
            time.sleep(PAUSE)
        total = sum(v for v in data[key].values())
        print(f"  Σ {total:,.0f}")

    # ИИ-источники — разбиваем на несколько запросов (лимит API 20 условий)
    ai_chunks = ai_filter_chunks()
    print(f"\n② ИИ-источники ({len(ai_chunks)} частей)  ", end="", flush=True)
    for y, m, ym in months:
        d1, d2 = month_range(y, m)
        for chunk_filt in ai_chunks:
            chunk_data = fetch_one_month(token, d1, d2, chunk_filt, "ym:s:visits")
            for path, val in chunk_data.items():
                data["ai"][(path, ym)] = data["ai"].get((path, ym), 0) + val
        print("·", end="", flush=True)
        time.sleep(PAUSE)
    total = sum(v for v in data["ai"].values())
    print(f"  Σ {total:,.0f}")

    # Собираем все страницы
    all_paths = set(titles.keys())
    for d in data.values():
        all_paths.update(p for p, _ in d.keys())
    all_paths = sorted(p for p in all_paths if "/articles" in p)
    print(f"\nВсего уникальных страниц: {len(all_paths)}")

    # xlsx
    print("\n⑤ Генерация xlsx…")
    build_xlsx(all_paths, titles, data, months, out_path)


if __name__ == "__main__":
    main()
